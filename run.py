"""
Script de entrada para Cloudera - Todo en uno
Compatible con entornos Jupyter/Notebook
"""
import os
from datetime import date, datetime
from flask import Flask, jsonify, send_file, send_from_directory
from flask_cors import CORS
from openpyxl import load_workbook

# Configuración
print("🚀 Iniciando Calendario YPF API (Flask) - Cloudera...")

# Cambiar al directorio del proyecto HermesCalendar
project_dir = '/home/cdsw/HermesCalendar'
if os.path.exists(project_dir):
    os.chdir(project_dir)
    print(f"📂 Cambiado a: {project_dir}")
else:
    print(f"⚠️ No se encontró {project_dir}, usando directorio actual")

current_dir = os.getcwd()
print(f"📂 Working Directory: {current_dir}")

EXCEL_PATH = os.path.join(current_dir, 'Calendario_Argentina_2026_Completo.xlsx')
print(f"📂 Excel: {EXCEL_PATH}")

# Crear app Flask
app = Flask(__name__)
CORS(app)

# Cache de datos
_cache = {
    "calendario": [],
    "hitos_periodicos": [],
    "hitos_importantes": [],
}


def cargar_excel():
    """Carga el Excel en memoria."""
    print(f"📂 Cargando: {EXCEL_PATH}")

    if not os.path.exists(EXCEL_PATH):
        print(f"❌ No se encontró: {EXCEL_PATH}")
        return False

    wb = load_workbook(EXCEL_PATH, read_only=True, data_only=True)

    # Calendario
    ws = wb["Calendario_Argentina_2026_Compl"]
    calendario = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            break
        fecha = row[0] if isinstance(row[0], datetime) else None
        if not fecha:
            continue
        calendario.append({
            "fecha": fecha.date().isoformat(),
            "nombreDia": row[1] or "",
            "tipo": row[2] or "No habil",
            "diaHabilNumero": row[3],
            "observaciones": row[4],
            "esHabil": row[2] == "Habil",
            "esFeriado": "feriado" in (row[4] or "").lower(),
        })

    # Hitos Periódicos
    ws = wb["Hitos_Periodicos"]
    hitos_per = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            break
        hitos_per.append({
            "id": f"per-{row[0]}",
            "diaHabil": row[0],
            "accion": row[1],
            "activo": True,
        })

    # Hitos Importantes
    ws = wb["Hitos_Importantes_YPF"]
    hitos_imp = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            break
        fecha = row[0] if isinstance(row[0], datetime) else None
        if not fecha:
            continue
        hoy = date.today()
        dias = (fecha.date() - hoy).days
        hitos_imp.append({
            "id": f"imp-{fecha.date().isoformat()}",
            "fecha": fecha.date().isoformat(),
            "titulo": row[1],
            "activo": True,
            "diasHastaHito": dias,
        })

    wb.close()

    _cache["calendario"] = calendario
    _cache["hitos_periodicos"] = hitos_per
    _cache["hitos_importantes"] = sorted(hitos_imp, key=lambda x: x["fecha"])

    print(f"✅ Cargado: {len(calendario)} días, {len(hitos_per)} hitos periódicos, {len(hitos_imp)} hitos importantes")
    return True


def obtener_posicion_actual():
    """Calcula D+X actual."""
    hoy = date.today()
    dia_actual = 1
    proximo = 2
    fecha_proximo = hoy.isoformat()

    for dia in _cache["calendario"]:
        dia_fecha = date.fromisoformat(dia["fecha"])
        if dia_fecha.year == hoy.year and dia_fecha.month == hoy.month:
            if dia_fecha <= hoy and dia["esHabil"]:
                dia_actual = dia["diaHabilNumero"] or dia_actual
            elif dia_fecha > hoy and dia["esHabil"]:
                proximo = dia["diaHabilNumero"] or proximo
                fecha_proximo = dia["fecha"]
                break

    return dia_actual, proximo, hoy.isoformat(), fecha_proximo


@app.route("/")
def home():
    """Servir el frontend HTML."""
    frontend_path = os.path.join(current_dir, "frontend", "app.html")
    if os.path.exists(frontend_path):
        return send_file(frontend_path)

    return jsonify({
        "status": "ok",
        "message": "Calendario YPF API (Flask)",
        "version": "1.0.0",
        "docs": "/api/v1/dashboard",
        "data_loaded": len(_cache["calendario"]) > 0,
    })


@app.route("/health")
def health():
    """Health check endpoint."""
    return jsonify({
        "status": "ok",
        "message": "Calendario YPF API (Flask)",
        "version": "1.0.0",
        "docs": "/api/v1/dashboard",
        "data_loaded": len(_cache["calendario"]) > 0,
    })


@app.route("/api/v1/dashboard")
def dashboard():
    """Dashboard principal."""
    if not _cache["calendario"]:
        if not cargar_excel():
            return jsonify({"error": "No se pudo cargar el Excel"}), 500

    dia_actual, proximo, fecha_actual, fecha_proximo = obtener_posicion_actual()
    hoy = date.today()

    hitos_proximo = [h for h in _cache["hitos_periodicos"] if h["diaHabil"] == proximo]
    feriados = [d for d in _cache["calendario"] if date.fromisoformat(d["fecha"]) >= hoy and d["esFeriado"]][:5]
    hitos_imp = [h for h in _cache["hitos_importantes"] if date.fromisoformat(h["fecha"]) >= hoy][:10]

    dias_restantes = len([
        d for d in _cache["calendario"]
        if date.fromisoformat(d["fecha"]).year == hoy.year
        and date.fromisoformat(d["fecha"]).month == hoy.month
        and date.fromisoformat(d["fecha"]) >= hoy
        and d["esHabil"]
    ])

    return jsonify({
        "posicionActual": {
            "diaHabilActual": dia_actual,
            "proximoDiaHabil": proximo,
            "fechaActual": fecha_actual,
            "fechaProximoDiaHabil": fecha_proximo,
        },
        "hitosProximoDia": {
            "diaHabil": proximo,
            "fecha": fecha_proximo,
            "hitosPeriodicos": hitos_proximo,
        },
        "feriadosProximos": feriados,
        "hitosImportantesProximos": hitos_imp,
        "diasHabilesRestantesMes": dias_restantes,
    })


@app.route("/api/v1/hitos-periodicos")
def hitos_periodicos():
    if not _cache["hitos_periodicos"]:
        cargar_excel()
    return jsonify(_cache["hitos_periodicos"])


@app.route("/api/v1/hitos-importantes")
def hitos_importantes():
    if not _cache["hitos_importantes"]:
        cargar_excel()
    return jsonify(_cache["hitos_importantes"])


@app.route("/api/v1/calendario")
def calendario():
    if not _cache["calendario"]:
        cargar_excel()
    return jsonify(_cache["calendario"])


@app.route("/api/v1/reload", methods=["POST"])
def reload():
    if cargar_excel():
        return jsonify({"status": "ok", "message": "Recargado"})
    return jsonify({"error": "Error al recargar"}), 500


@app.route("/frontend/<path:filename>")
def serve_frontend(filename):
    frontend_dir = os.path.join(current_dir, "frontend")
    return send_from_directory(frontend_dir, filename)


@app.route("/imagenes/<path:filename>")
def serve_images(filename):
    images_dir = os.path.join(current_dir, "imagenes")
    return send_from_directory(images_dir, filename)


# Cargar datos al inicio
cargar_excel()

# Obtener puerto
PORT = int(os.environ.get("CDSW_APP_PORT", os.environ.get("PORT", 8080)))

# Limpiar puerto si está ocupado (matar procesos zombies)
print(f"🔍 Limpiando puerto {PORT}...")
import subprocess
import time

try:
    # Intentar matar procesos en el puerto
    result = subprocess.run(
        f"lsof -ti:{PORT} | xargs -r kill -9",
        shell=True,
        capture_output=True,
        text=True
    )
    if result.returncode == 0:
        print(f"✅ Procesos en puerto {PORT} terminados")
        time.sleep(2)  # Esperar a que el SO libere el puerto
    else:
        print(f"✅ No hay procesos previos en puerto {PORT}")
except Exception as e:
    print(f"⚠️ Error limpiando puerto: {e}")

print(f"✅ Aplicación Flask lista")
print(f"🚀 Iniciando servidor en puerto {PORT}...")

# Reintentar bind si falla (race condition con Jupyter/Cloudera)
import sys
MAX_RETRIES = 3
for attempt in range(1, MAX_RETRIES + 1):
    try:
        print(f"📡 Intento {attempt}/{MAX_RETRIES}...")
        app.run(
            host="0.0.0.0",
            port=PORT,
            debug=False,
            threaded=True,
            use_reloader=False
        )
        break  # Si llega aquí, funcionó
    except OSError as e:
        if "Address already in use" in str(e) and attempt < MAX_RETRIES:
            print(f"⚠️ Puerto ocupado, esperando 3 segundos...")
            time.sleep(3)
            # Intentar matar de nuevo
            subprocess.run(f"lsof -ti:{PORT} | xargs -r kill -9", shell=True)
            time.sleep(2)
        else:
            print(f"❌ Error después de {attempt} intentos")
            raise
