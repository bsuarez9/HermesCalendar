"""
Punto de entrada SIMPLIFICADO de la aplicación FastAPI.
Solo lectura del Excel - Sin base de datos.
"""
import os
from datetime import date, datetime, timedelta
from typing import List, Optional

from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook

# Cargar variables de entorno
load_dotenv()

# Crear aplicación FastAPI
app = FastAPI(
    title="Calendario YPF API (Simple)",
    description="API REST para visualización del calendario de hitos de YPF - Solo lectura desde Excel",
    version="1.0.0-simple",
)

# Configurar CORS
origins = os.getenv("ALLOWED_ORIGINS", "http://localhost:5173,http://localhost:3000").split(",")
app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Path al Excel
EXCEL_PATH = os.getenv("EXCEL_FILE_PATH", "../Calendario_Argentina_2026_Completo.xlsx")

# Cache de datos en memoria
_cache = {
    "calendario": [],
    "hitos_periodicos": [],
    "hitos_importantes": [],
    "last_load": None,
}


def cargar_datos_excel():
    """Carga datos del Excel en memoria."""
    if not os.path.exists(EXCEL_PATH):
        raise FileNotFoundError(f"No se encontró el archivo Excel: {EXCEL_PATH}")

    print(f"📂 Cargando datos desde: {EXCEL_PATH}")

    wb = load_workbook(EXCEL_PATH, read_only=True, data_only=True)

    # Leer Calendario
    ws_calendario = wb["Calendario_Argentina_2026_Compl"]
    calendario = []
    for row in ws_calendario.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            break
        fecha = row[0] if isinstance(row[0], datetime) else None
        if not fecha:
            continue

        calendario.append({
            "fecha": fecha.date().isoformat(),
            "nombreDia": row[1] or "",
            "tipo": row[2] or "No habil",
            "diaHabilNumero": row[3],
            "observaciones": row[4],
            "esHabil": row[2] == "Habil",
            "esFeriado": "feriado" in (row[4] or "").lower(),
        })

    # Leer Hitos Periódicos
    ws_hitos_per = wb["Hitos_Periodicos"]
    hitos_periodicos = []
    for row in ws_hitos_per.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            break
        hitos_periodicos.append({
            "id": f"per-{row[0]}",
            "diaHabil": row[0],
            "accion": row[1],
            "activo": True,
        })

    # Leer Hitos Importantes
    ws_hitos_imp = wb["Hitos_Importantes_YPF"]
    hitos_importantes = []
    for row in ws_hitos_imp.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            break
        fecha = row[0] if isinstance(row[0], datetime) else None
        if not fecha:
            continue

        hoy = date.today()
        dias_hasta = (fecha.date() - hoy).days

        hitos_importantes.append({
            "id": f"imp-{fecha.date().isoformat()}-{row[1][:10]}",
            "fecha": fecha.date().isoformat(),
            "titulo": row[1],
            "activo": True,
            "diasHastaHito": dias_hasta,
        })

    wb.close()

    _cache["calendario"] = calendario
    _cache["hitos_periodicos"] = hitos_periodicos
    _cache["hitos_importantes"] = sorted(
        hitos_importantes, key=lambda x: x["fecha"]
    )
    _cache["last_load"] = datetime.now()

    print(f"✅ Datos cargados: {len(calendario)} días, {len(hitos_periodicos)} hitos periódicos, {len(hitos_importantes)} hitos importantes")


def obtener_dia_habil_actual():
    """Obtiene el día hábil actual (D+X)."""
    hoy = date.today()
    mes_actual = hoy.month
    anio_actual = hoy.year

    # Buscar en el calendario
    for dia in _cache["calendario"]:
        dia_fecha = date.fromisoformat(dia["fecha"])
        if (
            dia_fecha.year == anio_actual
            and dia_fecha.month == mes_actual
            and dia_fecha <= hoy
            and dia["esHabil"]
        ):
            dia_habil_actual = dia["diaHabilNumero"]

    # Buscar próximo día hábil
    proximo_dia_habil = None
    fecha_proximo = None
    for dia in _cache["calendario"]:
        dia_fecha = date.fromisoformat(dia["fecha"])
        if dia_fecha > hoy and dia["esHabil"]:
            proximo_dia_habil = dia["diaHabilNumero"]
            fecha_proximo = dia["fecha"]
            break

    return dia_habil_actual, proximo_dia_habil or dia_habil_actual + 1, hoy.isoformat(), fecha_proximo or hoy.isoformat()


@app.on_event("startup")
async def startup_event():
    """Cargar datos al iniciar."""
    cargar_datos_excel()


@app.get("/", tags=["Health"])
async def root():
    """Health check endpoint."""
    return {
        "status": "ok",
        "message": "Calendario YPF API (Simple) is running",
        "version": "1.0.0-simple",
        "docs": "/docs",
        "excel_loaded": _cache["last_load"] is not None,
    }


@app.get("/api/v1/dashboard")
async def obtener_dashboard():
    """Obtiene todos los datos del dashboard."""
    if not _cache["calendario"]:
        cargar_datos_excel()

    # Calcular posición actual
    dia_actual, proximo_dia, fecha_actual, fecha_proximo = obtener_dia_habil_actual()

    # Obtener hitos del próximo día
    hitos_proximo_dia = [
        h for h in _cache["hitos_periodicos"] if h["diaHabil"] == proximo_dia
    ]

    # Obtener próximos feriados
    hoy = date.today()
    feriados_proximos = [
        d for d in _cache["calendario"]
        if date.fromisoformat(d["fecha"]) >= hoy and d["esFeriado"]
    ][:5]

    # Obtener próximos hitos importantes
    hitos_importantes_proximos = [
        h for h in _cache["hitos_importantes"]
        if date.fromisoformat(h["fecha"]) >= hoy
    ][:10]

    # Días hábiles restantes en el mes
    mes_actual = hoy.month
    anio_actual = hoy.year
    dias_habiles_restantes = len([
        d for d in _cache["calendario"]
        if date.fromisoformat(d["fecha"]).year == anio_actual
        and date.fromisoformat(d["fecha"]).month == mes_actual
        and date.fromisoformat(d["fecha"]) >= hoy
        and d["esHabil"]
    ])

    return {
        "posicionActual": {
            "diaHabilActual": dia_actual,
            "proximoDiaHabil": proximo_dia,
            "fechaActual": fecha_actual,
            "fechaProximoDiaHabil": fecha_proximo,
        },
        "hitosProximoDia": {
            "diaHabil": proximo_dia,
            "fecha": fecha_proximo,
            "hitosPeriodicos": hitos_proximo_dia,
        },
        "feriadosProximos": feriados_proximos,
        "hitosImportantesProximos": hitos_importantes_proximos,
        "diasHabilesRestantesMes": dias_habiles_restantes,
    }


@app.get("/api/v1/hitos-periodicos")
async def listar_hitos_periodicos():
    """Lista todos los hitos periódicos."""
    if not _cache["hitos_periodicos"]:
        cargar_datos_excel()
    return _cache["hitos_periodicos"]


@app.get("/api/v1/hitos-importantes")
async def listar_hitos_importantes():
    """Lista todos los hitos importantes."""
    if not _cache["hitos_importantes"]:
        cargar_datos_excel()
    return _cache["hitos_importantes"]


@app.post("/api/v1/reload")
async def reload_excel():
    """Recarga los datos del Excel (útil después de editar el archivo)."""
    try:
        cargar_datos_excel()
        return {"status": "ok", "message": "Datos recargados exitosamente"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


if __name__ == "__main__":
    import uvicorn

    uvicorn.run("main_simple:app", host="0.0.0.0", port=8000, reload=True)
